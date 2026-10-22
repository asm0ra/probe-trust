"""Backend docx (IEEE, two columns)."""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from paper_refs import REF
OUT = "Probe-Trust_IEEE.docx"
ORDER = []
def C(*keys):
    nums = []
    for k in keys:
        if k not in ORDER: ORDER.append(k)
        nums.append(ORDER.index(k) + 1)
    return ", ".join(f"[{n}]" for n in nums)

# ---------------- document ----------------
d = Document(); sec = d.sections[0]; sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
sec.left_margin = sec.right_margin = Cm(1.6); sec.top_margin = Cm(1.9); sec.bottom_margin = Cm(2.5)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(3); st.paragraph_format.line_spacing = 1.0
def cols(section, n):
    sp = section._sectPr; c = sp.find(qn("w:cols"))
    if c is None: c = OxmlElement("w:cols"); sp.append(c)
    c.set(qn("w:num"), str(n)); c.set(qn("w:space"), "245")
def new_section(n): s = d.add_section(WD_SECTION.CONTINUOUS); cols(s, n); return s
def P(text="", bold=False, italic=False, size=None, align="j", indent=True, after=None, caps=False, before=None):
    p = d.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if caps: r.font.small_caps = True
    p.alignment = {"j": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT}[align]
    if indent and align == "j": p.paragraph_format.first_line_indent = Cm(0.4)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    return p
ROM = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV", "XVI", "XVII", "XVIII", "XIX", "XX", "XXI", "XXII"]; SECN = [0]; SUBN = [0]
def H1(t, num=True):
    if num: SECN[0] += 1; SUBN[0] = 0; return P(f"{ROM[SECN[0]-1]}. {t}", size=10, align="c", indent=False, caps=True, before=6)
    return P(t, size=10, align="c", indent=False, caps=True, before=6)
def H2(t): SUBN[0] += 1; return P(f"{chr(64+SUBN[0])}. {t}", italic=True, size=10, align="l", indent=False, before=3)
def mixed(parts, indent=True, size=None):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent: p.paragraph_format.first_line_indent = Cm(0.4)
    for t, b, i in parts:
        r = p.add_run(t); r.bold = b; r.italic = i
        if size: r.font.size = Pt(size)
    return p
def wide(fn): new_section(1); fn(); new_section(2)
FIGN = [0]; TABN = [0]; EQN = [0]
def FIG(path, cap, width=17.0):
    FIGN[0] += 1
    def f(): d.add_picture(path, width=Cm(width)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER; P(f"Fig. {FIGN[0]}. {cap}", size=8, align="c", indent=False)
    wide(f); return FIGN[0]
def TAB(cap, rows, widths=None, fs=7.5, bold_rows=()):
    TABN[0] += 1
    def f():
        P(f"TABLE {ROM[TABN[0]-1]}", size=8, align="c", indent=False, after=0, caps=True); P(cap, size=8, align="c", indent=False, caps=False, after=2)
        t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, c in enumerate(row):
                cell = t.cell(i, j); cell.text = ""; r = cell.paragraphs[0].add_run(str(c)); r.font.size = Pt(fs); r.bold = (i == 0 or i in bold_rows)
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        if widths:
            for row in t.rows:
                for j, w in enumerate(widths): row.cells[j].width = Cm(w)
        d.add_paragraph().paragraph_format.space_after = Pt(2)
    wide(f); return TABN[0]
def EQ(parts):  # parts: [(text, 'n'|'i'|'sub'|'sup')]
    EQN[0] += 1; p = d.add_paragraph(); pf = p.paragraph_format; pf.space_after = Pt(4); pf.space_before = Pt(2)
    pf.tab_stops.add_tab_stop(Cm(4.2), WD_TAB_ALIGNMENT.CENTER); pf.tab_stops.add_tab_stop(Cm(8.5), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    for t, k in parts:
        r = p.add_run(t); r.italic = (k == "i"); r.font.subscript = (k == "sub"); r.font.superscript = (k == "sup")
    p.add_run(f"\t({EQN[0]})"); return EQN[0]
def FRONT(title, authors, affil):
    cols(sec, 1)
    P(title, bold=True, size=18, align="c", indent=False, after=6)
    P(authors, size=11, align="c", indent=False, after=1)
    P(affil, italic=True, size=9, align="c", indent=False, after=6)
    new_section(2)
def FINISH():
    H1("References", num=False)
    for i, k in enumerate(ORDER, 1): P(f"[{i}] {REF[k]}", size=8, align="l", indent=False, after=1)
    unused = [k for k in REF if k not in ORDER]
    assert not unused, f"unused sources: {unused}"
    d.save(OUT); print("done", OUT, "references:", len(ORDER), "figures:", FIGN[0], "tables:", TABN[0])
